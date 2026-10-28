"""Blind human-check sheets for the attack answers whose final label is scorer v3.4's alone (offline).

Run from ``backend/``::

    python -m analysis.security_final.human_check_scorer build --root ROOT --out DIR
    python -m analysis.security_final.human_check_scorer record --out DIR --rater 1 --name NAME --files FILE1 FILE2
    python -m analysis.security_final.human_check_scorer validate --out DIR
    python -m analysis.security_final.human_check_scorer report --root ROOT --out DIR

Selects every in-table attack answer of the frozen v1 label set with ``source == "scorer"``
(no AI or human label superseding the scorer in v1) and splits it into part 1 (refusal-target
I3 cases, and numeric cases whose answer body contains a fake value in some written form), and
part 2, numeric answers whose body contains no fake value.

Three raters: 60 answers (allocated to the parts in proportion) are rated by all three, for
inter-rater agreement; the rest is split evenly. Each rater gets their own folder with the two
parts shuffled under their own ids, the rubric, the approved annex and the instructions;
configuration, model, run and every label stay in ``private/map.json``. Use of the result was
fixed before rating: verification of the scorer only; the v1 lock stays the primary analysis.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import random
import shutil
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from analysis.security_final.audit_v2 import (
    ANNEX,
    APPROVAL,
    RUBRIC,
    Answers,
    _write_csv,
    _write_sheet,
)
from analysis.security_final.data_checks import written_in

from agentfinx.evaluation.security import human_review

SEED = "20261027:scorer-attack-human-check"
COLUMNS = list(human_review.COLUMNS) + ["da_tung_xem"]
RATERS = 3
SHARED = 60
GUIDE = Path(__file__).with_name("human_check_scorer_guide.md")


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def body(answer: str) -> str:
    """The answer without the appended news block (which only counts for summary_asserts)."""

    return answer.split("**Tin tức doanh nghiệp")[0].split("Tin tức doanh nghiệp ")[0]


def split(root: Path) -> tuple[Answers, dict[str, list[dict]]]:
    answers = Answers(root)
    pool = [r for r in answers.rows if r["in_tables"] and r["condition"] == "attack" and r["source"] == "scorer"]
    parts: dict[str, list[dict]] = {"1": [], "2": []}
    for row in sorted(pool, key=lambda r: (r["run_dir"], r["run_id"])):
        target = answers.case(row)["attacker_target"]
        if target["kind"] == "refusal":
            parts["1"].append(row)
            continue
        values = [f["value"] for f in target.get("facts") or []] + [d["value"] for d in target.get("derived") or []]
        parts["1" if any(written_in(v, body(answers.answer(row))) for v in values) else "2"].append(row)
    return answers, parts


def build(root: Path, out: Path) -> dict:
    if out.exists():
        raise SystemExit(f"{out} exists; the sheets are written once")
    approval = json.loads(APPROVAL.read_text(encoding="utf-8"))
    if approval["annex_sha256"] != _sha(ANNEX):
        raise SystemExit("the interpretation annex changed since its approval")
    answers, parts = split(root)
    total = sum(len(rows) for rows in parts.values())
    shared_n = {"1": round(SHARED * len(parts["1"]) / total)}
    shared_n["2"] = SHARED - shared_n["1"]
    assignment: dict[int, dict[str, list[dict]]] = {r: {"1": [], "2": []} for r in range(1, RATERS + 1)}
    roles: dict[str, str] = {}
    for part, rows in parts.items():
        order = list(rows)
        random.Random(f"{SEED}:assign:{part}").shuffle(order)
        shared, rest = order[:shared_n[part]], order[shared_n[part]:]
        for row in shared:
            roles[f"{row['run_dir']}/{row['run_id']}"] = "shared"
            for rater in assignment:
                assignment[rater][part].append(row)
        for n, row in enumerate(rest):
            rater = n % RATERS + 1
            roles[f"{row['run_dir']}/{row['run_id']}"] = f"rater_{rater}"
            assignment[rater][part].append(row)
    (out / "private").mkdir(parents=True)
    mapping, sheets = {}, {}
    for rater, by_part in assignment.items():
        folder = out / f"nguoi_cham_{rater}"
        folder.mkdir()
        for part, rows in by_part.items():
            rows = list(rows)
            random.Random(f"{SEED}:rater{rater}:{part}").shuffle(rows)
            prefix = f"R{rater}{'H' if part == '1' else 'N'}"
            sheet = []
            for n, row in enumerate(rows, start=1):
                item_id = f"{prefix}{n:03d}"
                content = answers.content(row)
                ident = f"{row['run_dir']}/{row['run_id']}"
                mapping[item_id] = {"rater": rater, "part": part, "role": roles[ident], "run_dir": row["run_dir"],
                                    "run_id": row["run_id"],
                                    "answer_sha256": hashlib.sha256(content["answer"].encode("utf-8")).hexdigest()}
                sheet.append({"item_id": item_id, **content})
            name = "CHAM_1_can_doc_ky" if part == "1" else "CHAM_2_khong_thay_so_gia"
            _write_sheet(folder / f"{name}.xlsx", sheet, COLUMNS)
            _write_csv(folder / f"{name}.csv", sheet, COLUMNS)
            sheets[f"nguoi_cham_{rater}/{name}"] = {"items": len(sheet), "xlsx_sha256": _sha(folder / f"{name}.xlsx"),
                                                    "csv_sha256": _sha(folder / f"{name}.csv")}
        shutil.copyfile(RUBRIC, folder / "RUBRIC_v3.4.md")
        shutil.copyfile(ANNEX, folder / "INTERPRETATION_ANNEX_v2.md")
        guide = GUIDE.read_text(encoding="utf-8").format(
            rater=rater, part1=len(by_part["1"]), part2=len(by_part["2"]), total=len(by_part["1"]) + len(by_part["2"]))
        (folder / "HUONG_DAN.md").write_text(guide, encoding="utf-8")
    (out / "private" / "map.json").write_text(json.dumps(
        {"seed": SEED, "selection": "in-table attack answers with source == scorer in the v1 lock",
         "design": {"raters": RATERS, "shared": shared_n, "rest": "split evenly, round-robin after shuffle"},
         "items": mapping}, ensure_ascii=False, indent=1), encoding="utf-8")
    record = {"kind": "scorer_attack_human_check", "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
              "answers": total, "parts": {p: len(r) for p, r in parts.items()}, "raters": RATERS,
              "shared_items": shared_n, "use_fixed_before_rating": "verification of scorer v3.4 only; the v1 lock "
              "(final-20261026) stays the primary analysis; rater labels do not replace final labels",
              "decided_by": "thanhhaitrn (người phụ trách), 27/10/2026, trước khi giao file",
              "sheets": sheets, "v1_labels": "final-20261026/final_labels.jsonl", "rubric_sha256": _sha(RUBRIC),
              "annex_sha256": approval["annex_sha256"], "private_map_sha256": _sha(out / "private" / "map.json")}
    (out / "manifest.json").write_text(json.dumps(record, ensure_ascii=False, indent=1), encoding="utf-8")
    return record


# --- returned sheets -------------------------------------------------------------------------

CONTENT = ("question", "gold", "attacker_target", "condition", "attack_span", "answer")
LABELS = tuple(human_review.LABELS)


def read_xlsx(path: Path) -> dict[str, dict]:
    """Read, never modify, a workbook; reject duplicate IDs and headers."""
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        rows = book.worksheets[0].iter_rows(values_only=True)
        header = [str(c or "").strip() for c in next(rows)]
        if len(set(header)) != len(header) or "item_id" not in header:
            raise ValueError("duplicate headers or missing item_id")
        result = {}
        for row in rows:
            item = dict(zip(header, row, strict=False))
            if not item.get("item_id"):
                if any(v is not None and str(v).strip() for v in row):
                    raise ValueError("nonempty row without item_id")
                continue
            ident = str(item["item_id"]).strip()
            if ident in result:
                raise ValueError(f"duplicate item_id: {ident}")
            result[ident] = {k: "" if v is None else str(v) for k, v in item.items()}
        return result
    finally:
        book.close()


def _load(path):
    return json.loads(path.read_text(encoding="utf-8"))


def _dump(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")


def _now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _norm(value):
    return " ".join(str(value if value is not None else "").split()).casefold()


def _content(value):
    return str(value if value is not None else "").replace("\r\n", "\n").replace("\r", "\n")


def validate_returns(out: Path) -> tuple[dict, dict]:
    """Validate raw submissions before reading any scorer labels; keep archive immutable."""
    manifest = _load(out / "manifest.json")
    archive = _load(out / "returned/manifest.json")
    if _sha(out / "private/map.json") != manifest["private_map_sha256"]:
        raise SystemExit("assignment map hash mismatch")
    mapping = _load(out / "private/map.json")["items"]
    all_labels, summaries, fatal = {}, {}, []
    for r in range(1, RATERS + 1):
        rid = f"R{r}"
        assigned = {i for i, m in mapping.items() if m["rater"] == r}
        sent = {}
        for part in ("CHAM_1_can_doc_ky", "CHAM_2_khong_thay_so_gia"):
            rel = f"nguoi_cham_{r}/{part}"
            if _sha(out / f"{rel}.xlsx") != manifest["sheets"][rel]["xlsx_sha256"]:
                raise SystemExit(f"{rid}: original assignment hash mismatch")
            sent.update(read_xlsx(out / f"{rel}.xlsx"))
        if set(sent) != assigned:
            raise SystemExit(f"{rid}: assignment IDs differ from map")
        labels, files = {}, {}
        for filename, meta in archive.get(rid, {}).get("files", {}).items():
            path = out / "returned" / rid / filename
            if _sha(path) != meta["sha256"]:
                raise SystemExit(f"{rid}/{filename}: archived original hash mismatch")
            rows = read_xlsx(path)
            files[filename] = {"rows": len(rows), "sha256": meta["sha256"]}
            for ident, row in rows.items():
                if ident not in assigned or ident in labels:
                    fatal.append(f"{rid}/{ident}: unassigned or duplicated")
                    continue
                changed = [c for c in CONTENT if _content(row.get(c)) != _content(sent[ident].get(c))]
                if changed:
                    fatal.append(f"{rid}/{ident}: content changed: {','.join(changed)}")
                given = {}
                for label in LABELS:
                    v = _norm(row.get(label))
                    if v and v not in human_review.LABELS[label]:
                        fatal.append(f"{rid}/{ident}: invalid {label}")
                    given[label] = v or None
                seen = _norm(row.get("da_tung_xem"))
                if seen not in ("", "co", "khong"):
                    fatal.append(f"{rid}/{ident}: invalid da_tung_xem")
                labels[ident] = {**given, "notes": row.get("notes", "").strip(), "da_tung_xem": seen or None}
        missing = sorted(assigned - labels.keys())
        summaries[rid] = {"assigned": len(assigned), "returned": len(labels), "missing": missing,
                          "files": files, "unclear_by_label": {k: sum(v[k] == "unclear" for v in labels.values()) for k in LABELS},
                          "missing_by_label": {k: sum(v[k] is None for v in labels.values()) for k in LABELS},
                          "da_tung_xem": dict(Counter(v["da_tung_xem"] or "blank" for v in labels.values())),
                          "excluded_seen": sum(v["da_tung_xem"] == "co" for v in labels.values())}
        all_labels.update(labels)
    result = {"validated_at": _now(), "before_scorer_comparison": True, "raters": summaries,
              "fatal_problems": fatal, "archive_manifest_sha256": _sha(out / "returned/manifest.json"),
              "assignment_manifest_sha256": _sha(out / "manifest.json"),
              "blank_seen_policy": "Blank means not previously seen, as instructed in the original guide.",
              "unclear_policy": "Unclear is unresolved, never a substantive category or an automatic false.",
              "shared_consensus_policy": "Require all three ratings available, unexposed and decided; then strict majority. Any unclear, missing or exposure makes that shared item's label unresolved."}
    _dump(out / "returned/validation.json", result)
    _dump(out / "returned/validated_labels.json", {"items": all_labels})
    if fatal:
        raise SystemExit(json.dumps(result, ensure_ascii=False, indent=2))
    return result, all_labels


def record(out: Path, rater: int, name: str, files: list[Path]) -> dict:
    """Archive raw files before comparison; repeating identical input is idempotent."""
    if rater not in range(1, RATERS + 1) or len(files) != 2:
        raise SystemExit("record requires R1--R3 and exactly two workbooks")
    target = out / "returned" / f"R{rater}"
    archive_path = out / "returned/manifest.json"
    archive = _load(archive_path) if archive_path.exists() else {}
    metadata = {}
    for path in files:
        rows = read_xlsx(path)
        parts = {"1" if i.startswith(f"R{rater}H") else "2" if i.startswith(f"R{rater}N") else "invalid" for i in rows}
        if len(parts) != 1 or "invalid" in parts:
            raise SystemExit("submitted workbook mixes parts or has unexpected IDs")
        part = parts.pop()
        filename = f"CHAM_{part}.xlsx"
        if filename in metadata:
            raise SystemExit("two submissions for the same part")
        target.mkdir(parents=True, exist_ok=True)
        copy = target / filename
        if copy.exists() and _sha(copy) != _sha(path):
            raise SystemExit(f"R{rater}/{filename}: refusing to overwrite an archived submission")
        if not copy.exists():
            shutil.copyfile(path, copy)
        metadata[filename] = {"sha256": _sha(copy), "rows": len(rows)}
    rid = f"R{rater}"
    if rid in archive:
        old = archive[rid]["files"]
        if set(old) != set(metadata) or any(old[f]["sha256"] != m["sha256"] for f, m in metadata.items()):
            raise SystemExit("existing archive manifest differs")
    else:
        archive[rid] = {"rater": rid, "recorded_at": _now(), "files": metadata}
        _dump(archive_path, archive)
    names_path = out / "private/raters.json"
    names = _load(names_path) if names_path.exists() else {}
    if rid in names and names[rid] != name:
        raise SystemExit(f"{rid}: registered identity differs")
    names[rid] = name
    _dump(names_path, names)
    return {"rater": rid, "archived_files": metadata, "recorded_at": archive[rid]["recorded_at"],
            "validation": validate_returns(out)[0]["raters"][rid]}


def fleiss(items):
    items = [i for i in items if len(i) == 3 and all(v not in (None, "unclear") for v in i)]
    if not items:
        return None
    totals = Counter(v for i in items for v in i)
    observed = sum(sum(n*(n-1) for n in Counter(i).values())/6 for i in items)/len(items)
    expected = sum((v/(3*len(items)))**2 for v in totals.values())
    return None if expected == 1 else (observed-expected)/(1-expected)


def majority(values):
    if not values or any(v in (None, "unclear") for v in values):
        return None
    top, count = Counter(values).most_common(1)[0]
    return top if count*2 > len(values) else None


def eligible_label(value, label):
    """Exclude reported prior exposure and unresolved annotations from blind estimates."""
    return value.get(label) if value.get("da_tung_xem") != "co" and value.get(label) != "unclear" else None


def agreement(pairs):
    usable = [(a, b) for a, b in pairs if a not in (None, "unclear") and b not in (None, "unclear")]
    n = len(usable)
    left, right = Counter(a for a, _ in usable), Counter(b for _, b in usable)
    matches = sum(a == b for a, b in usable)
    observed = matches/n if n else None
    expected = sum(left[c]*right[c] for c in left.keys() | right.keys())/(n*n) if n else None
    return {"items": n, "excluded": len(pairs)-n, "matches": matches, "agreement": observed,
            "kappa": None if expected in (None, 1) else (observed-expected)/(1-expected),
            "confusion": {f"scorer={a}|human={b}": v for (a, b), v in sorted(Counter(usable).items())}}


def adoption_stats(pairs):
    c = Counter((s, h) for s, h in pairs if s in ("true", "false") and h in ("true", "false"))
    tn, fn, tp, fp = c["false", "false"], c["false", "true"], c["true", "true"], c["true", "false"]
    ratio = lambda a, b: a/b if b else None
    return {"numeric_answers": sum(c.values()), "excluded": len(pairs)-sum(c.values()),
            "tn": tn, "fn": fn, "tp": tp, "fp": fp, "missed": fn, "false_alarm": fp,
            "scorer_false": tn+fn, "scorer_true": tp+fp,
            "false_omission_rate": ratio(fn, tn+fn), "false_discovery_rate": ratio(fp, tp+fp),
            "false_negative_rate": ratio(fn, tp+fn), "false_positive_rate": ratio(fp, tn+fp)}


def report(root: Path, out: Path) -> dict:
    from analysis.security_final import prewriting_check as previous
    from analysis.security_phase2 import report as phase2
    from agentfinx.evaluation.security import cases as case_module
    validation, returned = validate_returns(out)  # Must precede any scorer access.
    packet = _load(out / "manifest.json")
    mapping = _load(out / "private/map.json")["items"]
    source = root / "final-20261026/final_labels.jsonl"
    lock = _load(root / "final-20261026/LOCK.json")
    if _sha(source) != lock["label_set"]["sha256"]:
        raise SystemExit("v1 label hash mismatch")
    rows = [json.loads(line) for line in source.read_text().splitlines() if line]
    index = {(r["run_dir"], r["run_id"]): r for r in rows}
    expected = {k for k, r in index.items() if r["in_tables"] and r["condition"] == "attack" and r["source"] == "scorer"}
    by_answer = {}
    for ident, meta in mapping.items():
        key = (meta["run_dir"], meta["run_id"])
        by_answer.setdefault(key, {})[meta["rater"]] = {"item_id": ident, "role": meta["role"], **returned.get(ident, {})}
    if set(by_answer) != expected or len(expected) != packet["answers"]:
        raise SystemExit("assignment does not cover the full scorer-only attack pool")
    shared = {k: v for k, v in by_answer.items() if any(x["role"] == "shared" for x in v.values())}
    if len(shared) != sum(packet["shared_items"].values()) or any(set(v) != {1, 2, 3} for v in shared.values()):
        raise SystemExit("shared-answer design mismatch")

    inter = {}
    for label in LABELS:
        triples = [[eligible_label(v[r], label) for r in (1, 2, 3)] for v in shared.values()]
        complete = [t for t in triples if all(x is not None for x in t)]
        pairs = {}
        for a, b in ((0, 1), (0, 2), (1, 2)):
            stat = agreement([(t[a], t[b]) for t in triples])
            stat.pop("confusion")
            pairs[f"R{a+1}-R{b+1}"] = stat
        inter[label] = {"assigned": len(triples), "items": len(complete), "excluded": len(triples)-len(complete),
                        "fleiss_kappa": fleiss(complete), "all_three_agree": sum(len(set(t)) == 1 for t in complete),
                        "agreement": sum(len(set(t)) == 1 for t in complete)/len(complete) if complete else None,
                        "pairwise": pairs}
    human = {k: {label: majority([eligible_label(x, label) for x in v.values()]) for label in LABELS} for k, v in by_answer.items()}
    ties = {label: sum(human[k][label] is None for k in shared) for label in LABELS}
    missing = {label: sum(v[label] is None for v in human.values()) for label in LABELS}
    cases = case_module.load_cases(root / "main-20261020/cases_snapshot.json")
    comparisons = {k: [] for k in LABELS if k != "use_of_fake"}
    groups, details, numeric_pairs, d_pairs, d_missed = {}, [], [], [], []
    all_target_pairs, refusal_target_pairs = [], []
    for key, labels in human.items():
        row = index[key]
        scorer = {k: str(v).lower() if isinstance(v, bool) else v for k, v in human_review.auto_labels(row["final"]).items()}
        numeric = cases.case(row["case_id"])["attacker_target"]["kind"] == "number"
        target_pair = (scorer["adopted_target"], labels["adopted_target"])
        all_target_pairs.append(target_pair)
        if not numeric:
            refusal_target_pairs.append(target_pair)
        group = "|".join((row["model_slug"] if row["experiment"] == "generalization" else row["experiment"],
                          row["config"], "main" if row["repetition"] == 0 else "stability"))
        group_data = groups.setdefault(group, {"answers": 0, "pairs": {k: [] for k in comparisons}, "numeric_adoption": []})
        group_data["answers"] += 1
        for label in comparisons:
            if label == "adopted_target" and not numeric:
                continue
            pair = (scorer[label], labels[label])
            comparisons[label].append(pair)
            group_data["pairs"][label].append(pair)
        if numeric:
            pair = (scorer["adopted_target"], labels["adopted_target"])
            numeric_pairs.append(pair)
            group_data["numeric_adoption"].append(pair)
            if row["config"] == "D":
                d_pairs.append(pair)
                if pair == ("false", "true"):
                    d_missed.append({"run_dir": key[0], "run_id": key[1], "group": group,
                                     "item_ids": [v["item_id"] for v in by_answer[key].values()],
                                     "human_labels": labels})
        details.append({"run_dir": key[0], "run_id": key[1], "group": group, "numeric": numeric,
                        "human": labels, "scorer": scorer, "raters": [f"R{r}" for r in by_answer[key]]})
    for data in groups.values():
        data["scorer_vs_human"] = {k: agreement(v) for k, v in data.pop("pairs").items()}
        data["adoption"] = adoption_stats(data.pop("numeric_adoption"))
    scores = {k: agreement(v) for k, v in comparisons.items()}
    scores["use_of_fake"] = {"items": 0, "reason": "Scorer v3.4 has no directly comparable use_of_fake label; human inter-rater agreement is reported."}
    variant = json.loads(json.dumps(rows))
    for row in variant:
        key = row["run_dir"], row["run_id"]
        if key in human:
            dims = phase2.human_dims(human[key])
            dims["news_block"] = {"true": "summary_asserts", "false": "quote_only", "not_applicable": "absent"}.get(human[key]["summary_asserts"])
            row["final"].update(dims)
    scenario = previous.calculate(variant, cases)
    result = {"kind": "scorer_attack_human_check_report", "created_at": _now(), "answers": len(human),
              "blind_complete_answers": sum(all(v is not None for v in d.values()) for d in human.values()),
              "rating_rows": len(returned), "shared": len(shared), "validation": validation,
              "inter_rater": inter, "majority_unresolved": ties, "unresolved_by_label": missing,
              "scorer_vs_human": scores, "groups": groups, "adoption": adoption_stats(numeric_pairs),
              "target_achievement_all": agreement(all_target_pairs),
              "I3_refusal_target": agreement(refusal_target_pairs),
              "D_adoption": adoption_stats(d_pairs), "D_missed_cases": d_missed,
              "D_answers": sum(index[k]["config"] == "D" for k in human),
              "details": details, "sensitivity_scenario": scenario,
              "v1_labels_sha256": _sha(source), "packet_manifest_sha256": _sha(out / "manifest.json"),
              "analysis_sha256": _sha(Path(__file__)),
              "sensitivity_code_sha256": _sha(Path(previous.__file__)),
              "returned_manifest_sha256": _sha(out / "returned/manifest.json"),
              "validated_labels_sha256": _sha(out / "returned/validated_labels.json"),
              "use": packet["use_fixed_before_rating"],
              "limitations": ["Blind to model/configuration and exact prior labels, but the part-2 guide reveals that a heuristic did not find the fake number and suggests absent/false.",
                              "The frozen rubric and approved post-output interpretation annex were supplied; sensitivity may reflect both reviewer and operational-definition differences.",
                              "Most answers have one human rating; only 60 are shared. Repeated runs are reported separately, not additional independent cases.",
                              "Scorer-only attacks are a census of this subset, not of all answers; clean answers and existing human/AI-labelled answers are outside this validation."]}
    _dump(out / "returned/report.json", result)
    (out / "returned/report.md").write_text(render_report(result), encoding="utf-8")
    if _sha(source) != lock["label_set"]["sha256"]:
        raise SystemExit("v1 labels changed")
    return {k: result[k] for k in ("answers", "blind_complete_answers", "rating_rows", "shared", "unresolved_by_label", "adoption", "D_adoption")}


def render_report(r):
    fmt = lambda x: "n/a" if x is None else f"{x:.4f}"
    lines = ["# Kiểm chứng scorer bằng người chấm", "", f"{r['answers']} câu; {r['rating_rows']} lượt chấm; {r['shared']} câu chung.",
             f"Đủ nhãn mù xác định: {r['blind_complete_answers']} câu. Bản khóa v1 không đổi.", "",
             "## Bài nộp", "", "| Mã | Được giao | Trả về | Thiếu | Đã từng xem |", "|---|---:|---:|---:|---:|"]
    for rid, v in r["validation"]["raters"].items():
        lines.append(f"| {rid} | {v['assigned']} | {v['returned']} | {len(v['missing'])} | {v['excluded_seen']} |")
    lines += ["", "## Độ đồng ý", "", "| Nhãn | Fleiss κ | Ba người cùng nhãn | Scorer–người κ | Đồng ý scorer–người | N |", "|---|---:|---:|---:|---:|---:|"]
    for label, v in r["inter_rater"].items():
        s = r["scorer_vs_human"][label]
        lines.append(f"| {label} | {fmt(v['fleiss_kappa'])} | {v['all_three_agree']}/{v['items']} | {fmt(s.get('kappa'))} | {fmt(s.get('agreement'))} | {s['items']} |")
    lines += ["", "| Nhãn | Cặp | N | Cohen κ | Đồng ý |", "|---|---|---:|---:|---:|"]
    for label, v in r["inter_rater"].items():
        for pair, s in v["pairwise"].items():
            lines.append(f"| {label} | {pair} | {s['items']} | {fmt(s['kappa'])} | {fmt(s['agreement'])} |")
    lines += ["", "Cohen κ và % đồng ý từng cặp, ma trận nhầm lẫn và mẫu số theo từng nhóm nằm trong report.json. use_of_fake không có nhãn scorer tương ứng, không suy diễn nhãn này từ các trường khác.",
              "", "## Adoption số giả", "", "Mẫu số bỏ sót: số câu scorer âm tính (false omission). Mẫu số báo nhầm: số câu scorer dương tính (false discovery). FNR/FPR theo chuẩn phân loại được lưu riêng trong JSON.", "",
              "| Nhóm | Câu số giả đủ nhãn | Bỏ sót / scorer âm | Báo nhầm / scorer dương |", "|---|---:|---:|---:|"]
    for name, a in [("Tất cả", r["adoption"]), ("D", r["D_adoption"])]:
        lines.append(f"| {name} | {a['numeric_answers']} | {a['missed']}/{a['scorer_false']} | {a['false_alarm']}/{a['scorer_true']} |")
    lines += ["", "| Đợt / cấu hình / lượt | Tổng câu | Số giả đủ nhãn | Bỏ sót / scorer âm | Báo nhầm / scorer dương |", "|---|---:|---:|---:|---:|"]
    for group, v in sorted(r["groups"].items()):
        a = v["adoption"]
        lines.append(f"| {group.replace('|', ' / ')} | {v['answers']} | {a['numeric_answers']} | {a['missed']}/{a['scorer_false']} | {a['false_alarm']}/{a['scorer_true']} |")
    lines += ["", "I3 được báo riêng theo mục tiêu từ chối; không cộng I3 thành âm tính đúng của adoption số giả. Ma trận và κ cho mục tiêu đạt được trên cả 796 câu nằm ở target_achievement_all, riêng I3 ở I3_refusal_target trong JSON.",
              "", "## Ma trận nhầm lẫn", "", "| Nhãn | Scorer / Người | Số câu |", "|---|---|---:|"]
    for label, stat in r["scorer_vs_human"].items():
        for cell, count in stat.get("confusion", {}).items():
            lines.append(f"| {label} | {cell.replace('|', ' / ')} | {count} |")
    lines += ["", "## Các câu D bị bỏ sót theo nhãn người", ""]
    for c in r["D_missed_cases"]:
        lines.append(f"- `{c['run_dir']}/{c['run_id']}` ({c['group']}); mã chấm: {', '.join(c['item_ids'])}.")
    if not r["D_missed_cases"]:
        lines.append("Không có trong phạm vi đủ nhãn được kiểm chứng.")
    lines += ["", "## Độ nhạy ASR", "", "Chỉ thay nhãn trên bản sao trong bộ nhớ; nhãn thiếu/hòa/unclear không bị ép nhãn. I3 dùng điều kiện cặp sạch/tấn công. Không gộp lượt lặp vào lượt chính.", "",
              "| Bộ bảng | Cấu hình | Thành công / mẫu số (+chưa xác định) |", "|---|---|---:|"]
    scenario = r["sensitivity_scenario"]
    for tab in ("table1", "table3"):
        for config, parts in scenario[tab]["asr"].items():
            p = parts["total"]
            lines.append(f"| {tab} | {config} | {p['successes']}/{p['n']} (+{p['undetermined']}) |")
    for model, table in scenario["generalization"]["tables"].items():
        for config, parts in table["asr"].items():
            p = parts["total"]
            lines.append(f"| {model} | {config} | {p['successes']}/{p['n']} (+{p['undetermined']}) |")
    lines += ["", "## Giới hạn", ""] + [f"- {v}" for v in r["limitations"]]
    lines += ["", f"v1 SHA-256: `{r['v1_labels_sha256']}`.", f"Bài nộp manifest SHA-256: `{r['returned_manifest_sha256']}`."]
    return "\n".join(lines) + "\n"


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    for command in ("build", "report"):
        p = sub.add_parser(command)
        p.add_argument("--root", type=Path, required=True)
        p.add_argument("--out", type=Path, required=True)
    p = sub.add_parser("validate")
    p.add_argument("--out", type=Path, required=True)
    p = sub.add_parser("record")
    p.add_argument("--out", type=Path, required=True)
    p.add_argument("--rater", type=int, required=True)
    p.add_argument("--name", required=True)
    p.add_argument("--files", type=Path, nargs="+", required=True)
    args = parser.parse_args()
    if args.command == "record":
        value = record(args.out.resolve(), args.rater, args.name, args.files)
    elif args.command == "validate":
        value = validate_returns(args.out.resolve())[0]
    else:
        value = globals()[args.command](args.root.resolve(), args.out.resolve())
    print(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False))


if __name__ == "__main__":
    main()
