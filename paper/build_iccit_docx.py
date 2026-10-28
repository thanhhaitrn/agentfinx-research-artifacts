"""Build an ICCIT-formatted .docx from main_iccit.tex (or main_iccit_vi.tex with LANG=vi).

Usage: python build_iccit_docx.py PAPER_DIR OUT.docx WORK_DIR [en|vi|en_long]
Run after latexmk, because labels and citation numbers are read from the .aux file.

1. Expand \\val numbers, refs and citations; turn inline math into plain text with sub/superscripts.
2. Replace every tabular with a placeholder and parse its rows; pandoc converts the prose.
3. python-docx sets Times New Roman 10 pt, 1-inch margins, running head, page numbers,
   writes the front matter, and builds the tables at the placeholders.
"""
import copy
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from docx.text.run import Run

paper, out, work = Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3])
LANG = sys.argv[4] if len(sys.argv) > 4 else "en"
work.mkdir(parents=True, exist_ok=True)
L = {
    "en": dict(
        stem="main_iccit", fig="fig_chain", fig_caption="Attack success along", table="Table", figure="Figure",
        intro="Introduction", refs="References", abstract="Abstract", keywords_label="Keywords",
        running_head="Robustness of a Financial Agent to News Attacks",
        title="Where Does a Financial-Report Agent\u2019s Robustness to News-Borne Attacks Come From?",
        authors=["[Author 1 full name]", "[Author 2 full name]", "[Author 3 full name]"],
        affil="[Department, University, full postal address, Country]",
        corresp="Corresponding author: [full name], [e-mail address]",
        keywords="Data poisoning; Financial question answering; Large language model agents; Prompt injection; "
                 "Retrieval-augmented generation"),
    "vi": dict(
        stem="main_iccit_vi", fig="fig_chain_vi", fig_caption="Tỉ lệ tấn công thành công dọc", table="Bảng",
        figure="Hình", intro="Giới thiệu", refs="Tài liệu tham khảo", abstract="Tóm tắt", keywords_label="Từ khoá",
        running_head="Tác tử tài chính trước tấn công qua tin tức",
        title="Khả năng chống tấn công qua tin tức của tác tử báo cáo tài chính đến từ đâu?",
        authors=["[Họ tên tác giả 1]", "[Họ tên tác giả 2]", "[Họ tên tác giả 3]"],
        affil="[Khoa, Trường, địa chỉ đầy đủ, Quốc gia]",
        corresp="Tác giả liên hệ: [họ tên], [địa chỉ e-mail]",
        keywords="Đầu độc dữ liệu; Hỏi đáp tài chính; Sinh văn bản tăng cường truy xuất; Tác tử mô hình ngôn ngữ "
                 "lớn; Tấn công chèn lệnh"),
    "en_long": dict(
        stem="main_iccit_long", table="Table", figure="Figure", intro="Introduction", refs="References",
        abstract="Abstract", keywords_label="Keywords",
        running_head="Robustness of a Financial Agent to News Attacks",
        title="Where Does a Financial-Report Agent\u2019s Robustness to News-Borne Attacks Come From?",
        author_runs=[("Tran Thanh Hai", "1"), ("Nguyen Tien Khue", "1"), ("Phan Trung Hai Nam", "1"),
                     ("Nguyen Khanh Van", "1"), ("Nguyen Duc Anh", "1"),
                     ("Nguyen Van Tang", "1,*")],
        affil_lines=[("1", "Faculty of Technology and Data Science, Foreign Trade University,"),
                     ("", "91 Chua Lang Street, Lang Thuong Ward, Hanoi, Vietnam")],
        corresp_line=("*", "Corresponding author: Nguyen Van Tang, tangnv@ftu.edu.vn"),
        keywords="Data poisoning; Financial question answering; Large language model agents; Prompt injection; "
                 "Retrieval-augmented generation"),
}[LANG]
RUNNING_HEAD = L["running_head"]
TITLE = L["title"]
FONT = "Times New Roman"


def balanced(s, i):
    assert s[i] == "{", s[i:i + 30]
    depth = 0
    for j in range(i, len(s)):
        if s[j] == "{":
            depth += 1
        elif s[j] == "}":
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
    raise ValueError("unbalanced")


def replace_cmd(s, name, fn, nargs=1):
    res, i, pat = [], 0, "\\" + name + "{"
    while True:
        k = s.find(pat, i)
        if k < 0:
            res.append(s[i:])
            return "".join(res)
        res.append(s[i:k])
        j, args = k + len(pat) - 1, []
        for _ in range(nargs):
            while s[j] in " \n":
                j += 1
            a, j = balanced(s, j)
            args.append(a)
        res.append(fn(*args))
        i = j


# ---------------------------------------------------------------- numbers, refs, cites
src = (paper / "generated/numbers.tex").read_text()
if (paper / "generated_extra/numbers_extra.tex").exists():
    src += (paper / "generated_extra/numbers_extra.tex").read_text()
nums, k = {}, 0
while (k := src.find("\\numdef{", k)) >= 0:
    key, j = balanced(src, k + 7)
    val, j = balanced(src, j)
    nums[key], k = val, j

aux = (paper / (L["stem"] + ".aux")).read_text()
labels = dict(re.findall(r"\\newlabel\{([^}]+)\}\{\{([^}]+)\}", aux))
cites = dict(re.findall(r"\\bibcite\{([^}]+)\}\{(\d+)\}", aux))

tex = (paper / (L["stem"] + ".tex")).read_text()
tex = re.sub(r"\{\\renewcommand\{\\small\}\{\\(?:script|footnote)size\}(?:\\setlength\{\\tabcolsep\}\{[^}]*\})?"
             r"(\\input\{[^}]+\})\}", r"\1", tex)
tex = re.sub(r"\\input\{(generated(?:_vi|_extra)?/tab_[^}]+)\}", lambda m: (paper / m.group(1)).read_text(), tex)
tex = replace_cmd(tex, "val", lambda key: nums[key])
tex = re.sub(r"\\ref\{([^}]+)\}", lambda m: labels[m.group(1)], tex)
tex = re.sub(r"~?\\cite\{([^}]+)\}",
             lambda m: " [" + ",".join(cites[c.strip()] for c in m.group(1).split(",")) + "]", tex)


def math_text(x):
    x = replace_cmd(x, "ensuremath", lambda y: y)
    x = x.replace("{\\times}", "\\times ")
    x = re.sub(r"_\{\\(?:mathrm|text)\{([^}]*)\}\}", r"\\textsubscript{\1}", x)
    x = re.sub(r"\^\{([^}]*)\}", lambda m: "\\textsuperscript{" + m.group(1).replace("-", "\u2212") + "}", x)
    for a, b in [("\\times", "\u00d7"), ("\\le", "\u2264"), ("\\to", "\u2192"), ("\\kappa", "\u03ba"),
                 ("^\\dagger", "\u2020"), ("\\dagger", "\u2020")]:
        x = x.replace(a, b)
    x = re.sub(r"(?<![A-Za-z\\{])([pnk])(?![A-Za-z])", r"\\emph{\1}", x)
    return x.replace(" ", "")


tex = replace_cmd(tex, "ensuremath", math_text)
tex = re.sub(r"\$([^$]+)\$", lambda m: math_text(m.group(1)), tex)
tex = replace_cmd(tex, "multirow", lambda a, b, c: c, nargs=3)
tex = re.sub(r"\\shortstack\[[lcr]\]\{", r"\\shortstack{", tex)
tex = replace_cmd(tex, "shortstack", lambda x: x.replace("\\\\", " "))

# ---------------------------------------------------------------- appendix headings and figures
if "\\appendix" in tex:
    head_part, app_part = tex.split("\\appendix", 1)
    app_part = re.sub(r"\\section\{([^}]*)\}\\label\{([^}]+)\}",
                      lambda m: "\\section*{Appendix " + labels[m.group(2)] + ". " + m.group(1) + "}", app_part)
    tex = head_part + app_part
tex = re.sub(r"\\resizebox\{[^}]*\}\{!\}\{\\input\{figures/([^}.]+)\.tex\}\}", r"\\includegraphics{figures/\1.png}", tex)
tex = re.sub(r"\\includegraphics(\[[^\]]*\])?\{figures/([^}.]+)\.(pdf|png)\}",
             lambda m: "\\includegraphics{" + str(paper / "figures" / (m.group(2) + ".png")) + "}", tex)


def number_figure(m):
    body = m.group(0)
    lab = re.search(r"\\label\{([^}]+)\}", body)
    if lab and lab.group(1) in labels:
        body = body.replace("\\caption{", "\\caption{\\textbf{" + L["figure"] + " " + labels[lab.group(1)] + ".} ", 1)
    return body


tex = re.sub(r"\\begin\{figure\}.*?\\end\{figure\}", number_figure, tex, flags=re.S)

# ---------------------------------------------------------------- tables -> placeholders
tables = {}


def parse_tabular(t):
    t = re.sub(r"\\(toprule|bottomrule)", "", t)
    rows, rules = [], set()
    for raw in re.split(r"\\\\(?:\[[^\]]*\])?", t):
        if "\\midrule" in raw:
            rules.add(len(rows))
            raw = raw.replace("\\midrule", "")
        raw = raw.strip()
        if not raw:
            continue
        cells = []
        for c in re.split(r"(?<!\\)&", raw):
            c = c.strip()
            m = re.match(r"\\multicolumn\{(\d+)\}\{[^}]*\}\{(.*)\}$", c, re.S)
            cells.append((m.group(2), int(m.group(1))) if m else (c, 1))
        rows.append(cells)
    return rows, rules


def fix_table(m):
    body = m.group(0)
    lab = re.search(r"\\label\{([^}]+)\}", body).group(1)
    cap, _ = balanced(body, body.index("\\caption{") + 8)
    parts = ["\n\n\\noindent TABLECAP " + labels[lab] + " " + cap + "\n\n"]
    pos, n = 0, 0
    while (start := body.find("\\begin{tabular}", pos)) >= 0:
        spec, after = balanced(body, start + len("\\begin{tabular}"))
        end = body.index("\\end{tabular}", after)
        tid = f"{lab}-{n}"
        rows, rules = parse_tabular(body[after:end])
        tables[tid] = (spec, rows, rules)
        parts.append(f"TABLEPH {tid}\n\n")
        pos, n = end + len("\\end{tabular}"), n + 1
    return "".join(parts)


tex = re.sub(r"\\begin\{table\*?\}.*?\\end\{table\*?\}", fix_table, tex, flags=re.S)

def fix_bib(m):
    b = m.group(1)
    b = re.sub(r"\\setlength\{[^}]*\}\{[^}]*\}|\\small", "", b)
    b = re.sub(r"\\bibitem\{([^}]+)\}", lambda mm: "\n\n{[}" + cites[mm.group(1)] + "{]} ", b)
    return "\\section*{" + L["refs"] + "}\n" + b


tex = re.sub(r"\\begin\{thebibliography\}\{[^}]*\}(.*?)\\end\{thebibliography\}", fix_bib, tex, flags=re.S)
head = "\\section{" + L["intro"] + "}"
body = head + tex.split(head, 1)[1]
(work / "body.tex").write_text("\\documentclass{article}\n\\begin{document}\n" + body)
ref = work / "ref.docx"
subprocess.run(["pandoc", "-o", str(ref), "--print-default-data-file", "reference.docx"], check=True)
subprocess.run(["pandoc", str(work / "body.tex"), "-f", "latex", "-t", "docx", "--reference-doc", str(ref),
                "--number-sections", "-o", str(work / "body.docx")], check=True)
doc = Document(work / "body.docx")

# ---------------------------------------------------------------- styles and page setup
for sec in doc.sections:
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.5)


def set_font(style, size=10, bold=None, italic=None, name=FONT):
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        rf.attrib.pop(qn(a), None)
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    style.font.color.rgb = None


for st in doc.styles:
    if st.type in (1, 2, 3):
        try:
            set_font(st)
        except Exception:
            pass
    if st.type == 1:
        pf = st.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(0), 1.0

for name, bold, italic, before, after in [("Heading 1", True, False, 10, 4), ("Heading 2", True, True, 8, 3),
                                          ("Heading 3", True, False, 6, 2)]:
    if name in doc.styles:
        st = doc.styles[name]
        set_font(st, 10, bold, italic)
        st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(before), Pt(after)
        st.paragraph_format.keep_with_next = True
for name in ("Body Text", "First Paragraph", "Compact"):
    if name in doc.styles:
        st = doc.styles[name]
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        st.paragraph_format.space_after = Pt(4)
if "Image Caption" in doc.styles:
    set_font(doc.styles["Image Caption"], 9, False, False)
    doc.styles["Image Caption"].paragraph_format.space_after = Pt(8)
if "Verbatim Char" in doc.styles:
    set_font(doc.styles["Verbatim Char"], 9, name="Courier New")

for p in doc.paragraphs:
    if p.style.name.startswith("Heading"):
        for r in p.runs:
            if "\t" in r.text:
                r.text = r.text.replace("\t", ". " if p.style.name == "Heading 1" else " ")


# ---------------------------------------------------------------- tables
def runs_from_latex(par, text, size):
    text = text.replace("---", "\u2014").replace("\\ldots", "\u2026").replace("\\,", "\u2009").replace("\\&", "&")
    text = text.replace("\\_", "_").replace("\\'{c}", "\u0107")
    text = text.replace("--", "\u2013").replace("\\%", "%").replace("\\ ", " ").replace("~", " ")
    text = text.replace("``", "\u201c").replace("''", "\u201d").replace("{,}", ",")
    text = text.replace("\\newline", " ")
    text = re.sub(r"\s+", " ", text).replace("`", "\u2018")
    text = re.sub(r"\\texttt\{([^{}]*)\}", r"\1", text)
    pos = 0
    for m in re.finditer(r"\\(emph|textsubscript|textsuperscript|textbf)\{([^{}]*)\}", text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        r = par.add_run(m.group(2))
        r.bold = m.group(1) == "textbf"
        r.italic = m.group(1) == "emph"
        r.font.subscript = m.group(1) == "textsubscript"
        r.font.superscript = m.group(1) == "textsuperscript"
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])
    for r in par.runs:
        r.font.name, r.font.size = FONT, Pt(size)


def plain(text):
    text = re.sub(r"\\[a-z]+\{([^{}]*)\}", r"\1", text)
    return text.replace("--", "-").replace("\\%", "%").replace("\\ ", " ")


def border(tc, edge, sz):
    tcPr = tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:color"), "000000")
    b.append(el)


def cell_margins(table, lr=0.04):
    mar = OxmlElement("w:tblCellMar")
    for edge in ("left", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:w"), str(int(lr * 1440)))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    table._tbl.tblPr.append(mar)


def build_table(anchor, spec, rows, rules):
    aligns, fixed = [], {}
    for m in re.finditer(r"(?:>\{(?:[^{}]|\{[^{}]*\})*\})?([lrcLp])(\{[^{}]*\})?", spec):
        kind, arg = m.group(1), m.group(2) or ""
        if kind in "Lp":
            frac = re.search(r"([\d.]+)\\textwidth", arg)
            fixed[len(aligns)] = 6.5 * float(frac.group(1)) if frac else 1.5
            aligns.append("l")
        else:
            aligns.append(kind)
    ncols = len(aligns)
    size = 8 if ncols <= 9 else 7
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell_margins(table)
    longest = [2] * ncols
    for ri, row in enumerate(rows):
        ci = 0
        for text, span in row:
            if ci >= ncols:
                break
            cell = table.cell(ri, ci)
            if span > 1:
                cell = cell.merge(table.cell(ri, ci + span - 1))
            par = cell.paragraphs[0]
            a = aligns[ci] if span == 1 else "c"
            par.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT,
                             "c": WD_ALIGN_PARAGRAPH.CENTER}[a]
            runs_from_latex(par, text, size)
            if span == 1:
                longest[ci] = max(longest[ci], len(plain(text)))
            ci += span
    char_w = size / 72 * 0.56
    widths = [fixed.get(i, n * char_w + 0.16) for i, n in enumerate(longest)]
    if sum(widths) > 6.5:
        widths = [w * 6.5 / sum(widths) for w in widths]
    for gc, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 1440)))
    last = len(rows) - 1
    for ri in range(len(rows)):
        tr = table.rows[ri]._tr
        col = 0
        for tc in tr.findall(qn("w:tc")):
            span_el = tc.find(qn("w:tcPr") + "/" + qn("w:gridSpan"))
            span = int(span_el.get(qn("w:val"))) if span_el is not None else 1
            tcPr = tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.insert(0, tcW)
            tcW.set(qn("w:w"), str(int(sum(widths[col:col + span]) * 1440)))
            tcW.set(qn("w:type"), "dxa")
            if ri == 0:
                border(tc, "top", 8)
            if ri in rules:
                border(tc, "top", 4)
            if ri == last:
                border(tc, "bottom", 8)
            col += span
        trPr = tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        if ri < last:
            for p in table.rows[ri].cells[0].paragraphs:
                p.paragraph_format.keep_with_next = True
    anchor._p.addnext(table._tbl)
    return table


for p in list(doc.paragraphs):
    t = p.text.strip()
    if t.startswith("TABLECAP "):
        num = t.split()[1]
        for r in p.runs:
            r.text = r.text.replace("TABLECAP " + num + " ", "", 1)
        newr = copy.deepcopy(p.runs[0]._r)
        p.runs[0]._r.addprevious(newr)
        nr = Run(newr, p)
        nr.text, nr.bold, nr.italic = f"{L['table']} {num}. ", True, False
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(8), Pt(4)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.first_line_indent = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for r in p.runs:
            r.font.size = Pt(9)
    elif t.startswith("TABLEPH "):
        tid = t.split()[1]
        spec, rows, rules = tables[tid]
        build_table(p, spec, rows, rules)
        # placeholder paragraph becomes an empty spacer placed after the table
        for r in p.runs:
            r.text = ""
        p.paragraph_format.space_after = Pt(4 if tid.endswith("-0") and f"{tid[:-2]}-1" in tables else 8)
        tbl = p._p.getnext()
        tbl.addnext(p._p)

# ---------------------------------------------------------------- front matter
first = doc.paragraphs[0]


def add_before(runs, align=WD_ALIGN_PARAGRAPH.CENTER, size=10, after=0):
    p = first.insert_paragraph_before()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    for text, fmt in runs:
        r = p.add_run(text)
        r.font.name, r.font.size = FONT, Pt(fmt.get("size", size))
        r.bold, r.italic = fmt.get("bold", False), fmt.get("italic", False)
        r.font.superscript = fmt.get("sup", False)
    return p


add_before([(TITLE, {"bold": True, "size": 12})], after=10)
author_runs = L.get("author_runs") or list(zip(L["authors"], ["1,*", "1", "2"]))
affil_lines = L.get("affil_lines") or [("1", L["affil"]), ("2", L["affil"])]
corresp_line = L.get("corresp_line") or ("*", L["corresp"])
runs = []
for i, (name, sup) in enumerate(author_runs):
    runs += [((", " if i else "") + name, {}), (sup, {"sup": True})]
add_before(runs, after=6)
for sup, text in affil_lines:
    add_before([(sup, {"sup": True, "size": 9}), (text, {"size": 9})])
add_before([(corresp_line[0], {"sup": True, "size": 9}), (corresp_line[1], {"size": 9})], after=6)
add_before([("Running head: ", {"italic": True, "size": 9}), (RUNNING_HEAD, {"size": 9})], after=10)
abstract = re.search(r"\\noindent\\textbf\{" + L["abstract"] + r"\.\}\n(.*?)\n\n", tex, re.S).group(1)
abstract = re.sub(r"\s+", " ", abstract).replace("'", "\u2019").strip()
add_before([(L["abstract"] + ". ", {"bold": True}), (abstract, {})], align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6)
add_before([(L["keywords_label"] + ": ", {"bold": True}), (L["keywords"], {})], align=WD_ALIGN_PARAGRAPH.LEFT, after=6)

# ---------------------------------------------------------------- header / footer
sec = doc.sections[0]
sec.different_first_page_header_footer = True
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = hp.add_run(RUNNING_HEAD)
hr.italic, hr.font.size, hr.font.name = True, Pt(9), FONT
for footer in (sec.footer, sec.first_page_footer):
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run()
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        r._r.append(el)
    r.font.size, r.font.name = Pt(9), FONT

for shape in doc.inline_shapes:
    ratio = shape.height / shape.width
    width = Inches(6.5) if ratio < 0.55 else Inches(4.2)  # wide diagrams use the full text width
    shape.width = width
    shape.height = int(width * ratio)

doc.core_properties.title = TITLE
doc.core_properties.author = ""
doc.save(out)
print("saved", out, "tables:", len(tables))
